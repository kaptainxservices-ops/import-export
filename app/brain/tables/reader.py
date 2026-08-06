"""Getting a grid of cells out of HTML bodies and spreadsheet files.

Reading only. No interpretation happens here — that is `columns` and `parser` — so
that a change to how a supplier's file is shaped never turns into a change to how
offers are understood.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from bs4 import BeautifulSoup

# Below this a 'table' is layout: a logo, a signature block, a two-cell header banner.
MIN_TABLE_ROWS = 4


@dataclass
class Grid:
    """A rectangular block of text cells, and where it came from."""

    rows: list[list[str]] = field(default_factory=list)
    source: str = ""

    def __len__(self) -> int:
        return len(self.rows)

    @property
    def width(self) -> int:
        return max((len(r) for r in self.rows), default=0)


def read_html_tables(html: str | None, min_rows: int = MIN_TABLE_ROWS) -> list[Grid]:
    """Every table in an HTML body that is big enough to be a price list.

    Only leaf tables are taken — a table that itself contains a table big enough to be
    a price list is a layout wrapper, not data.

    This is the opposite of the obvious rule, and getting it the wrong way round is
    quietly destructive rather than loud. Yukatel and Landotech nest tables several
    deep for formatting; keeping the outermost one yields cells whose text is the
    flattened content of everything inside, so a row reads
    'Dear Yogesh, I hope this message finds you well' in all six columns. That parses
    as a product with no error anywhere.
    """
    if not html:
        return []

    try:
        soup = BeautifulSoup(html, "lxml")
    except Exception:
        return []

    grids: list[Grid] = []
    for index, table in enumerate(soup.find_all("table")):
        if any(len(inner.find_all("tr")) >= min_rows for inner in table.find_all("table")):
            continue

        rows = []
        for tr in table.find_all("tr"):
            cells = [_clean(td.get_text(" ", strip=True)) for td in tr.find_all(["td", "th"])]
            if any(cells):
                rows.append(cells)

        if len(rows) >= min_rows:
            grids.append(Grid(rows=rows, source=f"table {index + 1}"))

    return grids


def read_spreadsheet(data: bytes, filename: str = "") -> list[Grid]:
    """Read .xlsx / .xlsm or legacy .xls into one grid per sheet.

    Legacy .xls needs a different library entirely, and one supplier in the sample
    (Parktel) still sends it. Failing to read a file is reported as an empty result
    rather than an exception — one unreadable attachment must not abandon the email.
    """
    name = filename.lower()
    if name.endswith(".xls"):
        return _read_xls(data, filename)
    return _read_xlsx(data, filename)


def _read_xlsx(data: bytes, filename: str) -> list[Grid]:
    try:
        import openpyxl
    except ImportError:  # pragma: no cover
        return []

    try:
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return []

    grids: list[Grid] = []
    try:
        for sheet_name in book.sheetnames:
            sheet = book[sheet_name]
            rows = [
                [_clean(_stringify(cell)) for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            rows = [r for r in rows if any(r)]
            if rows:
                grids.append(Grid(rows=rows, source=f"{filename or 'workbook'}!{sheet_name}"))
    finally:
        book.close()

    return grids


def _read_xls(data: bytes, filename: str) -> list[Grid]:
    try:
        import xlrd
    except ImportError:  # pragma: no cover
        return []

    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception:
        return []

    grids: list[Grid] = []
    for sheet in book.sheets():
        rows = []
        for index in range(sheet.nrows):
            cells = [_clean(_stringify(v)) for v in sheet.row_values(index)]
            if any(cells):
                rows.append(cells)
        if rows:
            grids.append(Grid(rows=rows, source=f"{filename or 'workbook'}!{sheet.name}"))

    return grids


def _stringify(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        # Excel stores every number as a float, so a quantity of 50 arrives as 50.0 and
        # an EAN as 4.25201190776e+12. Integral values are rendered as integers so that
        # '50' does not become '50.0' and fail to parse as a count.
        if value.is_integer():
            return str(int(value))
        return repr(value)
    return str(value)


_WHITESPACE = re.compile(r"[\s ​]+")


def _clean(text: str) -> str:
    """Collapse whitespace, including the non-breaking spaces Excel exports leave."""
    if not text:
        return ""
    return _WHITESPACE.sub(" ", str(text)).strip()
